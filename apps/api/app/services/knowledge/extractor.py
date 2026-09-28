import os
from abc import ABC, abstractmethod
from typing import BinaryIO
import PyPDF2
import docx
import csv

class ExtractedDocument:
    def __init__(self, content: str):
        self.content = content

class DocumentExtractor(ABC):
    @abstractmethod
    def supports(self, file_type: str) -> bool:
        pass

    @abstractmethod
    def extract(self, file: BinaryIO) -> ExtractedDocument:
        pass

class PDFExtractor(DocumentExtractor):
    def supports(self, file_type: str) -> bool:
        return file_type.lower() == "pdf"

    def extract(self, file: BinaryIO) -> ExtractedDocument:
        reader = PyPDF2.PdfReader(file)
        text = ""
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n\n"
        return ExtractedDocument(content=text)

class DocxExtractor(DocumentExtractor):
    def supports(self, file_type: str) -> bool:
        return file_type.lower() in ["docx", "doc"]

    def extract(self, file: BinaryIO) -> ExtractedDocument:
        doc = docx.Document(file)
        text = "\n".join([paragraph.text for paragraph in doc.paragraphs])
        return ExtractedDocument(content=text)

class TxtExtractor(DocumentExtractor):
    def supports(self, file_type: str) -> bool:
        return file_type.lower() == "txt"

    def extract(self, file: BinaryIO) -> ExtractedDocument:
        text = file.read().decode("utf-8", errors="replace")
        return ExtractedDocument(content=text)

class CsvExtractor(DocumentExtractor):
    def supports(self, file_type: str) -> bool:
        return file_type.lower() == "csv"

    def extract(self, file: BinaryIO) -> ExtractedDocument:
        content = file.read().decode("utf-8", errors="replace")
        lines = content.splitlines()
        reader = csv.reader(lines)
        text = ""
        for row in reader:
            text += ", ".join(row) + "\n"
        return ExtractedDocument(content=text)

class ExtractorFactory:
    def __init__(self):
        self.extractors = [
            PDFExtractor(),
            DocxExtractor(),
            TxtExtractor(),
            CsvExtractor()
        ]

    def get_extractor(self, file_type: str) -> DocumentExtractor:
        for ext in self.extractors:
            if ext.supports(file_type):
                return ext
        raise ValueError(f"Unsupported file type: {file_type}")
